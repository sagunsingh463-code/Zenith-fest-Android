import os
import sys
from pathlib import Path
import openpyxl

from backend.database import get_db_connection, init_db

# Ensure utf-8 encoding for console printing of emojis and symbols on Windows
if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events() -> None:
    """
    Seeds the events table from data/events_seed.xlsx.
    Skips if events are already seeded.
    Raises FileNotFoundError if data/events_seed.xlsx is missing.
    """
    init_db()

    if not SEED_FILE.exists():
        raise FileNotFoundError(
            f"Seed file not found at '{SEED_FILE}'. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if table already has rows to avoid duplicate seeding
    cursor.execute("SELECT COUNT(*) FROM events")
    existing_count = cursor.fetchone()[0]
    if existing_count > 0:
        print(f"Events table already contains {existing_count} rows. Skipping seed.")
        conn.close()
        return

    # Load Excel workbook
    wb = openpyxl.load_workbook(SEED_FILE)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError("Sheet 'Events' not found in events_seed.xlsx.")

    sheet = wb["Events"]

    # Verify header row
    expected_headers = [
        "id", "title", "category", "description",
        "event_date", "event_time", "venue", "capacity", "fee", "icon"
    ]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        conn.close()
        raise ValueError("events_seed.xlsx has no data.")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    for idx, expected in enumerate(expected_headers):
        if idx >= len(header_row) or header_row[idx] != expected:
            conn.close()
            raise ValueError(
                f"Header mismatch in row 1: expected '{expected}' at column {idx + 1}, found '{header_row[idx] if idx < len(header_row) else None}'"
            )

    # Insert rows
    inserted = 0
    for row in rows[1:]:
        if not row or all(cell is None for cell in row):
            continue

        (
            event_id, title, category, description,
            event_date, event_time, venue, capacity, fee, icon
        ) = row

        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(event_id) if event_id is not None else None,
                str(title).strip() if title else "",
                str(category).strip() if category else "",
                str(description).strip() if description else "",
                str(event_date).strip() if event_date else "",
                str(event_time).strip() if event_time else "",
                str(venue).strip() if venue else "",
                int(capacity) if capacity is not None else 0,
                int(fee) if fee is not None else 0,
                str(icon).strip() if icon else ""
            ),
        )
        inserted += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted} events into fest.db.")


if __name__ == "__main__":
    seed_events()

    # Display contents of events table
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    rows = cursor.fetchall()
    print(f"\n--- Events Table Content ({len(rows)} rows) ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} | {r['category']} | {r['event_date']} {r['event_time']} | {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()
